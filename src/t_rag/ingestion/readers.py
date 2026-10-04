from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from pathlib import Path

from t_rag.config import Settings, SourceConfig
from t_rag.ingestion.metadata import build_metadata


@dataclass(slots=True)
class Section:
    text: str
    location: str


@dataclass(slots=True)
class LoadedDocument:
    source: str
    relative_path: str
    text: str
    metadata: dict
    file_hash: str
    size_bytes: int
    summary_only: bool = False


def read_sections(path: Path) -> list[Section]:
    suffix = path.suffix.lower()
    if suffix in {".md", ".mdc", ".txt"}:
        text = path.read_text(encoding="utf-8-sig")
        if suffix == ".txt":
            return [Section(text, "全文")] if text.strip() else []
        text = re.sub(r"\A---[^\S\n]*\n.*?\n---[^\S\n]*\n", "", text, count=1, flags=re.S)
        sections, headings, lines = [], [], []
        fenced = False
        fence_char = ""
        fence_length = 0
        def flush():
            content = "\n".join(lines).strip()
            if content:
                sections.append(Section(content, " / ".join(h for _, h in headings) or "正文"))
            lines.clear()
        for line in text.splitlines():
            marker = re.match(r"^\s{0,3}(`{3,}|~{3,})", line)
            if marker:
                value = marker.group(1)
                if not fenced:
                    fenced, fence_char, fence_length = True, value[0], len(value)
                elif value[0] == fence_char and len(value) >= fence_length:
                    fenced = False
                lines.append(line)
                continue
            heading = None if fenced else re.match(r"^\s{0,3}(#{1,6})\s+(.+?)\s*#*\s*$", line)
            if heading:
                flush()
                level = len(heading.group(1))
                while headings and headings[-1][0] >= level:
                    headings.pop()
                headings.append((level, heading.group(2)))
            lines.append(line)
        flush()
        return sections
    if suffix == ".pdf":
        from pypdf import PdfReader
        result = []
        for number, page in enumerate(PdfReader(str(path)).pages, 1):
            text = page.extract_text() or ""
            if text.strip():
                result.append(Section(text, f"第 {number} 页"))
        if not result:
            raise ValueError(f"PDF 无可提取文本（首版不支持 OCR）：{path.name}")
        return result
    if suffix == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as stream:
            return _table_sections(csv.reader(stream), "CSV")
    if suffix == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            return [section for sheet in workbook.worksheets
                    for section in _table_sections(sheet.iter_rows(values_only=True), sheet.title)]
        finally:
            workbook.close()
    raise ValueError(f"不支持的文档格式：{suffix}")


def _table_sections(rows, sheet: str) -> list[Section]:
    result, group, first = [], [], 1
    header = ""
    last = 0
    for number, row in enumerate(rows, 1):
        last = number
        cells = ["" if cell is None else str(cell) for cell in row]
        if not any(cells):
            continue
        text = " | ".join(cells)
        if not header:
            header = text
        if not group:
            first = number
        group.append(f"行 {number}: {text}")
        if len(group) == 30:
            result.append(Section(f"表头: {header}\n" + "\n".join(group), f"{sheet} · 行 {first}–{number}"))
            group = []
    if group:
        result.append(Section(f"表头: {header}\n" + "\n".join(group), f"{sheet} · 行 {first}–{last}"))
    return result


def load_document(path: Path, relative_path: Path, source: SourceConfig, settings: Settings,
                  file_hash: str, size_bytes: int) -> LoadedDocument:
    sections = read_sections(path)
    text = "\n\n".join(s.text for s in sections)
    return LoadedDocument(source.name, relative_path.as_posix(), text,
                          build_metadata(source, relative_path, text), file_hash, size_bytes)

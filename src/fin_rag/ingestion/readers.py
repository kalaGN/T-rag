from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from pathlib import Path

from fin_rag.config import Settings, SourceConfig
from fin_rag.ingestion.metadata import build_metadata

_FRONTMATTER_PATTERN = re.compile(r"^---\s*\n(.*?)\n---\s*\n", re.DOTALL)


@dataclass(slots=True)
class LoadedDocument:
    source: str
    relative_path: str
    text: str
    metadata: dict[str, str | None]
    file_hash: str
    size_bytes: int
    summary_only: bool = False


def load_document(
    path: Path,
    relative_path: Path,
    source: SourceConfig,
    settings: Settings,
    file_hash: str,
    size_bytes: int,
) -> LoadedDocument:
    suffix = path.suffix.lower()
    if suffix in {".md", ".mdc", ".txt", ".py", ".java", ".xml"}:
        text = _read_text(path)
    elif suffix == ".csv":
        text = _read_csv(path, settings.large_file_threshold_bytes)
    elif suffix == ".xlsx":
        text = _read_xlsx(path)
    elif suffix == ".pdf":
        text = _read_pdf(path)
    else:
        text = f"[unsupported file type] {path.name}"
    metadata = build_metadata(source=source, relative_path=relative_path, text=text)
    summary_only = path.stat().st_size > settings.large_file_threshold_bytes and suffix in {".csv", ".xlsx"}
    return LoadedDocument(
        source=source.name,
        relative_path=relative_path.as_posix(),
        text=text,
        metadata=metadata,
        file_hash=file_hash,
        size_bytes=size_bytes,
        summary_only=summary_only,
    )


def _read_text(path: Path) -> str:
    try:
        raw = path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        import logging

        logger = logging.getLogger(__name__)
        logger.warning("File %s is not valid UTF-8, retrying with errors=ignore", path)
        raw = path.read_text(encoding="utf-8", errors="ignore")
    if path.suffix.lower() != ".mdc":
        return raw
    match = _FRONTMATTER_PATTERN.match(raw)
    if match:
        frontmatter = match.group(1).strip()
        body = raw[match.end():].strip()
        return f"[frontmatter]\n{frontmatter}\n\n{body}"
    return raw


def _read_csv(path: Path, large_file_threshold_bytes: int) -> str:
    if path.stat().st_size > large_file_threshold_bytes:
        with path.open("r", encoding="utf-8", errors="ignore", newline="") as file:
            reader = csv.reader(file)
            rows = list(_take(reader, 6))
        return _format_tabular_summary(path.name, rows)
    with path.open("r", encoding="utf-8", errors="ignore", newline="") as file:
        reader = csv.reader(file)
        rows = list(reader)
    return "\n".join(",".join(cell for cell in row) for row in rows)


def _read_xlsx(path: Path) -> str:
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError:
        return f"[xlsx dependency missing] {path.name}"

    workbook = load_workbook(path, read_only=True, data_only=True)
    parts: list[str] = []
    for sheet in workbook.worksheets:
        parts.append(f"[sheet] {sheet.title}")
        for row in _take(sheet.iter_rows(values_only=True), 10):
            values = ["" if cell is None else str(cell) for cell in row]
            parts.append(",".join(values))
    return "\n".join(parts)


def _read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ModuleNotFoundError:
        return f"[pdf dependency missing] {path.name}"

    reader = PdfReader(str(path))
    parts: list[str] = []
    for page_index, page in enumerate(reader.pages[:10], start=1):
        parts.append(f"[page {page_index}]")
        parts.append(page.extract_text() or "")
    return "\n".join(parts)


def _format_tabular_summary(filename: str, rows: list[list[str]]) -> str:
    header = rows[0] if rows else []
    sample_rows = rows[1:] if len(rows) > 1 else []
    lines = [f"[summary only] {filename}", f"columns: {header}"]
    for row in sample_rows:
        lines.append(f"sample: {row}")
    return "\n".join(lines)


def _take(iterator, size: int):
    count = 0
    for item in iterator:
        if count >= size:
            break
        yield item
        count += 1
